#!/usr/bin/env python
"""Turn the cleaned stage HTML reports into a designed PowerPoint deck.

The deck mirrors the visual language of the HTML reports: an accent colour
system, a consistent title band + footer on every slide, accent section
dividers, banded data tables, methodology call-outs and framed figures.

For each stage s01-s07 it emits a section divider, content slides (section
text / notes / tables) and one slide per generated figure.

Usage::

    python report_to_pptx.py --archive <fdh_dir> --out <deck.pptx> [--mode full|summary]
"""
from __future__ import annotations

import argparse
import io
import re
from pathlib import Path

from bs4 import BeautifulSoup
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

# --------------------------------------------------------------------------- #
# design system
# --------------------------------------------------------------------------- #
ACCENT = RGBColor(0x2C, 0x6E, 0x9B)
ACCENT_DK = RGBColor(0x1E, 0x4D, 0x6E)
ACCENT2 = RGBColor(0xC0, 0x84, 0x3B)
INK = RGBColor(0x23, 0x2B, 0x33)
MUT = RGBColor(0x6B, 0x77, 0x85)
LIGHT = RGBColor(0xEC, 0xF1, 0xF5)
BAND = RGBColor(0xF3, 0xF6, 0xF9)
LINE = RGBColor(0xD7, 0xDE, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xDC, 0xE8, 0xF1)
FONT = "Arial"

SW, SH = Inches(13.333), Inches(7.5)
CT_TOP = Inches(1.55)
CT_H = Inches(5.25)

STAGE_TITLES = {
    "s01": "Input preprocessing",
    "s02": "Homolog search",
    "s03": "MSA & conservation",
    "s04": "Complex (co-folding)",
    "s05": "Docking",
    "s06": "Interaction model",
    "s07": "Mutation generation",
}
STAGE_ROLE = {
    "s01": "Target sequence, ligands and methods-grade input QC",
    "s02": "Homolog mining and sequence-diversity clustering",
    "s03": "Multiple-sequence alignment and per-position conservation",
    "s04": "WT protein-ligand complex co-folding ensemble",
    "s05": "Reference docking and pose accuracy",
    "s06": "Ensemble interaction model and multi-engine pose classification",
    "s07": "Designable positions and mutation-candidate generation",
}
FIGURE_CAPTIONS = {
    "s01_sequence_composition": "Amino-acid composition of the target sequence.",
    "s02_identity_histogram": "Sequence identity to the target across retained homologs.",
    "s02_cluster_sizes": "Size distribution of the sequence-diversity clusters.",
    "s03_conservation_profile": "Per-position evolutionary conservation.",
    "s03_entropy_profile": "Per-position Shannon entropy of the MSA.",
    "s03_coverage_profile": "Per-position alignment coverage (1 - gap frequency).",
    "s04_ensemble_confidence": "Co-folding confidence metrics across the diffusion ensemble.",
    "s04_pae_heatmap": "Predicted aligned error (PAE) of the WT complex.",
    "s04_plddt_profile": "Per-residue model confidence (pLDDT) of the WT complex.",
    "s05_pose_rmsd": "Docking pose accuracy (RMSD to reference).",
    "s06_contact_frequency": "WT-ensemble pocket contact frequency by residue.",
    "s06_pose_reliability": "Reliability-metric distributions across the pose ensemble.",
    "s06_multi_engine_roles": "Multi-engine docking pose classification by engine.",
    "s07_candidates_by_generator": "Candidate count by generation strategy.",
    "s07_mutations_per_candidate": "Distribution of mutation order (single vs multipoint).",
    "s07_design_space_heatmap": "Normalised design-space metrics across designable positions.",
}
SKIP_SECTIONS = ("references",)
NUM_RE = re.compile(r"^[\-+]?[\d.,%/x±– ]+$")


# --------------------------------------------------------------------------- #
# low-level styling helpers
# --------------------------------------------------------------------------- #
def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _rect(slide, l, t, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, l, t, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s


def _box(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    return tf


def _run(p, text, size, color, bold=False, italic=False, font=FONT):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    return r


# --------------------------------------------------------------------------- #
# deck
# --------------------------------------------------------------------------- #
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = SW
        self.prs.slide_height = SH

    def _blank(self, bg=WHITE):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        _bg(s, bg)
        return s

    def _footer(self, s, tag):
        _rect(s, Inches(0.55), Inches(7.0), Inches(12.23), Pt(0.75), fill=LINE)
        tf = _box(s, Inches(0.55), Inches(7.06), Inches(9), Inches(0.32))
        _run(tf.paragraphs[0], "FDH · computational design pipeline", 8.5, MUT)
        tf2 = _box(s, Inches(9.8), Inches(7.06), Inches(2.98), Inches(0.32))
        p = tf2.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        _run(p, f"{tag}    {len(self.prs.slides):02d}", 8.5, MUT)

    def _content(self, eyebrow, title, tag):
        s = self._blank()
        tf = _box(s, Inches(0.6), Inches(0.34), Inches(12.1), Inches(0.32))
        _run(tf.paragraphs[0], eyebrow.upper(), 11.5, ACCENT, bold=True)
        tf2 = _box(s, Inches(0.57), Inches(0.62), Inches(12.2), Inches(0.78))
        _run(tf2.paragraphs[0], title, 25, INK, bold=True)
        _rect(s, Inches(0.6), Inches(1.4), Inches(12.13), Pt(2.4), fill=ACCENT)
        self._footer(s, tag)
        return s

    # -- slide types -------------------------------------------------------- #
    def cover(self, title, subtitle):
        s = self._blank()
        _rect(s, 0, 0, Inches(0.36), SH, fill=ACCENT)
        _rect(s, Inches(0.36), 0, Inches(0.06), SH, fill=ACCENT2)
        tf = _box(s, Inches(1.0), Inches(2.45), Inches(11.4), Inches(0.45))
        _run(tf.paragraphs[0], "COMPUTATIONAL ENZYME DESIGN", 13, ACCENT, bold=True)
        tf2 = _box(s, Inches(0.97), Inches(2.92), Inches(11.6), Inches(1.7))
        _run(tf2.paragraphs[0], title, 37, INK, bold=True)
        _rect(s, Inches(1.02), Inches(4.55), Inches(3.1), Pt(3), fill=ACCENT2)
        tf3 = _box(s, Inches(1.0), Inches(4.8), Inches(11.4), Inches(1.0))
        _run(tf3.paragraphs[0], subtitle, 15.5, MUT)

    def divider(self, key, title, role):
        s = self._blank(bg=ACCENT)
        _rect(s, Inches(0.85), Inches(2.55), Inches(1.7), Pt(4), fill=ACCENT2)
        tf = _box(s, Inches(0.82), Inches(2.7), Inches(11.6), Inches(1.0))
        _run(tf.paragraphs[0], key.upper(), 44, WHITE, bold=True)
        tf2 = _box(s, Inches(0.85), Inches(3.95), Inches(11.6), Inches(0.9))
        _run(tf2.paragraphs[0], title, 30, WHITE, bold=True)
        tf3 = _box(s, Inches(0.87), Inches(4.85), Inches(11.4), Inches(0.8))
        _run(tf3.paragraphs[0], role, 15.5, PALE, italic=True)

    def bullets(self, eyebrow, title, items, tag):
        s = self._content(eyebrow, title, tag)
        tf = _box(s, Inches(0.62), CT_TOP, Inches(12.1), CT_H)
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        first = True
        for text, kind in items:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(7)
            p.line_spacing = 1.04
            if kind == "head":
                p.space_before = Pt(7)
                _run(p, text, 15, ACCENT_DK, bold=True)
            elif kind == "note":
                _run(p, "◆  ", 11, ACCENT2, bold=True)
                _run(p, text, 11.5, MUT, italic=True)
            elif kind == "sub":
                p.level = 1
                _run(p, "–  ", 12.5, MUT)
                _run(p, text, 12.5, INK)
            else:
                _run(p, "▪  ", 13, ACCENT, bold=True)
                _run(p, text, 13, INK)

    def table(self, eyebrow, title, rows, tag, maxr=12, maxc=8):
        s = self._content(eyebrow, title, tag)
        trunc = len(rows) > maxr or max(len(r) for r in rows) > maxc
        rows = [r[:maxc] for r in rows[:maxr]]
        nr, nc = len(rows), max(len(r) for r in rows)
        rh = min(0.42, 5.0 / nr)
        gf = s.shapes.add_table(nr, nc, Inches(0.62), CT_TOP, Inches(12.1),
                                Inches(rh * nr))
        t = gf.table
        t.first_row = True
        cw = int(Inches(12.1) / nc)
        for col in t.columns:
            col.width = cw
        for i, row in enumerate(rows):
            t.rows[i].height = Inches(rh)
            for j in range(nc):
                cell = t.cell(i, j)
                cell.text = row[j] if j < len(row) else ""
                cell.margin_left = Inches(0.09)
                cell.margin_right = Inches(0.06)
                cell.margin_top = Pt(2)
                cell.margin_bottom = Pt(2)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                para = cell.text_frame.paragraphs[0]
                if i == 0:
                    cell.fill.fore_color.rgb = ACCENT
                    _style_cell(para, 10, WHITE, bold=True)
                else:
                    cell.fill.fore_color.rgb = WHITE if i % 2 else BAND
                    numeric = bool(para.runs and NUM_RE.match(para.runs[0].text.strip()))
                    _style_cell(para, 9.5, INK,
                                align=PP_ALIGN.RIGHT if (numeric and j > 0) else PP_ALIGN.LEFT)
        if trunc:
            tf = _box(s, Inches(0.62), Inches(6.62), Inches(12.1), Inches(0.34))
            _run(tf.paragraphs[0], "Showing a representative excerpt — full table in data.xlsx.",
                 10, MUT, italic=True)

    def figure(self, eyebrow, title, tif, caption, tag):
        s = self._content(eyebrow, title, tag)
        im = Image.open(tif).convert("RGB")
        w, h = im.size
        buf = io.BytesIO()
        im.save(buf, "PNG")
        buf.seek(0)
        maxw, maxh = Inches(11.6), Inches(4.7)
        ar = w / h
        width = maxw
        height = Emu(int(width / ar))
        if height > maxh:
            height = maxh
            width = Emu(int(height * ar))
        left = Emu(int((SW - width) / 2))
        pic = s.shapes.add_picture(buf, left, Inches(1.62), width=width, height=height)
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
        tf = _box(s, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, caption, 12, MUT, italic=True)

    def save(self, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        return len(self.prs.slides)


def _style_cell(para, size, color, bold=False, align=PP_ALIGN.LEFT):
    para.alignment = align
    if not para.runs:
        para.add_run().text = ""
    for r in para.runs:
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = FONT


# --------------------------------------------------------------------------- #
# HTML -> ordered content blocks
# --------------------------------------------------------------------------- #
def parse_report(html: str):
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    blocks = []

    def walk(node):
        for child in getattr(node, "children", []):
            name = getattr(child, "name", None)
            if name is None:
                continue
            if name in ("h1", "h2", "h3", "h4"):
                txt = child.get_text(" ", strip=True)
                if txt:
                    blocks.append((name, txt))
            elif name == "p":
                txt = child.get_text(" ", strip=True)
                if txt:
                    blocks.append(("p", txt))
            elif name in ("ul", "ol"):
                for li in child.find_all("li", recursive=False):
                    txt = li.get_text(" ", strip=True)
                    if txt:
                        blocks.append(("li", txt))
            elif name == "table":
                rows = []
                for tr in child.find_all("tr"):
                    cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
                    if cells:
                        rows.append(cells)
                if rows:
                    blocks.append(("table", rows))
            elif name == "div" and "warn" in (child.get("class") or []):
                txt = child.get_text(" ", strip=True)
                if txt:
                    blocks.append(("note", txt))
            elif name in ("div", "section", "footer", "header", "main", "article", "details"):
                walk(child)
    walk(soup.body or soup)
    return blocks


def add_content(deck, eyebrow, tag, blocks, mode):
    cur = ""
    items = []
    skip = False

    def flush():
        nonlocal items
        if not items:
            return
        budget, chunk = 0, []
        for it in items:
            ln = len(it[0])
            if chunk and (budget + ln > 1050 or len(chunk) >= 11):
                deck.bullets(eyebrow, cur or "Overview", chunk, tag)
                chunk, budget = [], 0
            chunk.append(it)
            budget += ln
        if chunk:
            deck.bullets(eyebrow, cur or "Overview", chunk, tag)
        items = []

    for kind, val in blocks:
        if kind in ("h1", "h2"):
            flush()
            cur = val
            skip = any(k in val.lower() for k in SKIP_SECTIONS)
        elif skip:
            continue
        elif kind == "h3":
            items.append((val, "head"))
        elif kind in ("p", "li", "note"):
            if mode != "full":
                continue
            items.append((val, "note" if kind == "note" else ("sub" if kind == "li" else "bullet")))
        elif kind == "table":
            flush()
            deck.table(eyebrow, cur or "Data", val, tag)
    flush()


def build(archive: Path, out: Path, mode: str):
    deck = Deck()
    title = "FDH"
    readme = archive / "README.md"
    if readme.exists():
        m = re.search(r"^#\s+(.+)$", readme.read_text(), re.M)
        if m:
            title = m.group(1).split(" - ")[0].strip()
    deck.cover(f"{title} — computational design pipeline",
               "Stages s01–s07  ·  figures, data and methods  ·  derived from the cleaned stage reports")

    for key, stage_title in STAGE_TITLES.items():
        sdir = archive / key
        rpt = sdir / "report.html"
        if not rpt.exists():
            continue
        eyebrow = f"{key} · {stage_title}"
        deck.divider(key, stage_title, STAGE_ROLE.get(key, ""))
        add_content(deck, eyebrow, key.upper(), parse_report(rpt.read_text(errors="ignore")), mode)
        for tif in sorted((sdir / "figure").glob("*.tif")):
            cap = FIGURE_CAPTIONS.get(tif.stem, "")
            fig_title = tif.stem.split("_", 1)[-1].replace("_", " ").capitalize()
            deck.figure(eyebrow, fig_title, tif, cap, key.upper())

    n = deck.save(out)
    print(f"WROTE {out}  ({n} slides)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--archive", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--mode", default="full", choices=["full", "summary"])
    args = ap.parse_args()
    build(args.archive, args.out, args.mode)


if __name__ == "__main__":
    main()
